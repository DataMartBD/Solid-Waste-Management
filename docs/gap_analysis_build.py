"""Render the KCC gap analysis as a Word document.

    server/.venv/Scripts/python.exe docs/gap_analysis_build.py

The findings themselves live in `gap_analysis_data.py`, shared with the PDF
renderer, so the two documents cannot drift apart. Edit them there.
"""

from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Inches

from gap_analysis_data import (
    CROSS_CUTTING,
    FINDINGS,
    PRIORITIES,
    first_sentence,
)

OUT = Path(__file__).resolve().parent / "SWMS-Gap-Analysis.docx"

INK = RGBColor(0x1A, 0x1A, 0x1A)
MUTED = RGBColor(0x5A, 0x5A, 0x5A)
GREEN = RGBColor(0x1B, 0x7F, 0x3B)
AMBER = RGBColor(0x9A, 0x5B, 0x00)
RED = RGBColor(0xB0, 0x2A, 0x2A)

STATUS_COLOUR = {"Built": GREEN, "Partial": AMBER, "Missing": RED}
STATUS_FILL = {"Built": "E7F4EA", "Partial": "FDF1E0", "Missing": "FBE9E9"}


# --------------------------------------------------------------------------- #
# Document construction
# --------------------------------------------------------------------------- #

def shade(cell, hex_fill):
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")
    el.set(qn("w:fill"), hex_fill)
    cell._tc.get_or_add_tcPr().append(el)


def run(paragraph, text, *, bold=False, size=10.5, colour=INK, italic=False):
    r = paragraph.add_run(text)
    r.bold = bold
    r.italic = italic
    r.font.size = Pt(size)
    r.font.color.rgb = colour
    r.font.name = "Calibri"
    return r


def heading(doc, text, level):
    sizes = {1: 19, 2: 14, 3: 11.5}
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18 if level < 3 else 12)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    run(p, text, bold=True, size=sizes[level])
    return p


def body(doc, text, *, size=10.5, colour=INK, space_after=8, italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    run(p, text, size=size, colour=colour, italic=italic)
    return p


def build():
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    for section in doc.sections:
        section.left_margin = section.right_margin = Inches(0.9)
        section.top_margin = section.bottom_margin = Inches(0.8)

    # ---- title page block ------------------------------------------------ #
    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(2)
    run(title, "Smart Sweep SWMS", bold=True, size=26)
    sub = doc.add_paragraph()
    sub.paragraph_format.space_after = Pt(14)
    run(sub, "Requirements gap analysis against the KCC stakeholder document",
        size=13, colour=MUTED)

    meta = doc.add_paragraph()
    meta.paragraph_format.space_after = Pt(16)
    run(meta, f"Prepared {date.today():%d %B %Y}   ·   "
              f"Source: Documents.docx (KCC, primary collectors, data management)   ·   "
              f"Reviewed against the live application",
        size=9.5, colour=MUTED)

    # ---- how to read ----------------------------------------------------- #
    heading(doc, "How to read this", 2)
    body(doc,
         "Every requirement in the source document has been checked against the "
         "application as it actually stands — the database models, the API routes and "
         "the screens — rather than against a feature list. Each item is marked:")

    legend = doc.add_table(rows=1, cols=3)
    legend.alignment = WD_TABLE_ALIGNMENT.LEFT
    for i, (label, meaning) in enumerate([
        ("Built", "In the product and usable today"),
        ("Partial", "The foundation exists; something specific is missing"),
        ("Missing", "Not addressed at all"),
    ]):
        cell = legend.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        run(p, f"{label}\n", bold=True, size=10, colour=STATUS_COLOUR[label])
        run(p, meaning, size=9.5, colour=MUTED)
        shade(cell, STATUS_FILL[label])
    body(doc, "", space_after=4)

    body(doc,
         "The headline: the household and property registry, billing and money trail, "
         "complaints, routes, vehicles and the new survey module are substantially "
         "built. The gaps cluster in three places — anything that has to leave the "
         "system and reach a person (notifications, campaigns, published tariffs), "
         "anything about the workforce beyond who they are (attendance, training, PPE, "
         "health, salary), and anything about service geography finer than a ward "
         "(road-wise operator assignment, dumping points, bins).")

    # ---- summary table --------------------------------------------------- #
    heading(doc, "Summary", 2)
    counts = {"Built": 0, "Partial": 0, "Missing": 0}
    for finding in FINDINGS:
        counts[finding[2]] += 1
    for item in CROSS_CUTTING:
        counts[item[1]] += 1
    total = sum(counts.values())
    body(doc,
         f"{total} requirement areas assessed: {counts['Built']} built, "
         f"{counts['Partial']} partial, {counts['Missing']} missing.")

    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    widths = [Inches(3.4), Inches(0.9), Inches(2.4)]
    header = table.rows[0].cells
    for i, label in enumerate(["Requirement", "Status", "Principal gap"]):
        header[i].text = ""
        run(header[i].paragraphs[0], label, bold=True, size=9.5)
        shade(header[i], "EFEFEF")
        header[i].width = widths[i]

    # Every row the counts above are drawn from, so the table and the sentence
    # introducing it can never disagree.
    rows = [(f[1], f[2], f[4]) for f in FINDINGS]
    rows += [(f"Cross-cutting: {c[0]}", c[1], c[2]) for c in CROSS_CUTTING]

    for requirement, status, missing in rows:
        cells = table.add_row().cells
        cells[0].text = ""
        run(cells[0].paragraphs[0], requirement, size=9.5)
        cells[1].text = ""
        run(cells[1].paragraphs[0], status, bold=True, size=9.5,
            colour=STATUS_COLOUR[status])
        shade(cells[1], STATUS_FILL[status])
        cells[2].text = ""
        short = first_sentence(missing)
        run(cells[2].paragraphs[0], short, size=9, colour=MUTED)
        for i, w in enumerate(widths):
            cells[i].width = w

    doc.add_page_break()

    # ---- detail ---------------------------------------------------------- #
    heading(doc, "Findings in detail", 1)
    current_section = None
    for section, requirement, status, have, missing in FINDINGS:
        if section != current_section:
            heading(doc, section, 2)
            current_section = section

        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run(p, requirement, bold=True, size=11.5)
        run(p, "   ")
        run(p, status.upper(), bold=True, size=9, colour=STATUS_COLOUR[status])

        lead = doc.add_paragraph()
        lead.paragraph_format.space_after = Pt(4)
        lead.paragraph_format.left_indent = Inches(0.15)
        run(lead, "What exists.  ", bold=True, size=10)
        run(lead, have, size=10)

        if missing.strip() != "—":
            gap = doc.add_paragraph()
            gap.paragraph_format.space_after = Pt(6)
            gap.paragraph_format.left_indent = Inches(0.15)
            run(gap, "What is missing.  ", bold=True, size=10, colour=RED)
            run(gap, missing, size=10)

    # ---- cross-cutting --------------------------------------------------- #
    heading(doc, "Cross-cutting requirements", 2)
    body(doc,
         "The document closes by asking for a system that is mobile-friendly, "
         "Bengali-supported, role-based and introduced gradually to workers with "
         "limited digital skills. Those are assessed separately here because they cut "
         "across every module.")

    for item in CROSS_CUTTING:
        name, status, note = item[0], item[1], item[2]
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        run(p, name, bold=True, size=11)
        run(p, "   ")
        run(p, status.upper(), bold=True, size=9, colour=STATUS_COLOUR[status])
        detail = doc.add_paragraph()
        detail.paragraph_format.left_indent = Inches(0.15)
        detail.paragraph_format.space_after = Pt(4)
        run(detail, note, size=10)

    doc.add_page_break()

    # ---- priorities ------------------------------------------------------ #
    heading(doc, "Recommended order of work", 1)
    body(doc,
         "Ordered by what the document treats as urgent, weighted towards work that "
         "unblocks several gaps at once and that builds on what is already there. "
         "Effort is relative to the modules already delivered, not an estimate in days.")

    for rank, name, refs, why, effort in PRIORITIES:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(11)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        run(p, f"{rank}.  ", bold=True, size=11.5, colour=MUTED)
        run(p, name, bold=True, size=11.5)
        run(p, f"   ({refs})", size=9.5, colour=MUTED)

        detail = doc.add_paragraph()
        detail.paragraph_format.left_indent = Inches(0.25)
        detail.paragraph_format.space_after = Pt(3)
        run(detail, why, size=10)

        eff = doc.add_paragraph()
        eff.paragraph_format.left_indent = Inches(0.25)
        eff.paragraph_format.space_after = Pt(2)
        run(eff, "Effort:  ", bold=True, size=9.5, colour=MUTED)
        run(eff, effort, size=9.5, colour=MUTED)

    # ---- closing note ---------------------------------------------------- #
    heading(doc, "A note on what is already strong", 2)
    body(doc,
         "The document’s stated first phase is “household registry, road-wise operator "
         "mapping, daily work reporting, complaint management, fee and royalty tracking, "
         "workforce and van records, and a KCC monitoring dashboard”. Six of those seven "
         "are built or close to it. Road-wise operator mapping is the one that is not "
         "started, and workforce records are the one that is thinnest — which is why "
         "they sit first and fourth in the list above.")
    body(doc,
         "The survey module delivered most recently also covers the document’s "
         "“survey data collection sheet” requirement more thoroughly than asked: the "
         "questionnaire is stored as data rather than code, so a revised form is a "
         "configuration change, and a reviewed survey can be promoted straight into the "
         "holding register.")

    footer = doc.add_paragraph()
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.space_before = Pt(24)
    run(footer,
        "Prepared from the source requirements document and a direct review of the "
        "application code.", size=8.5, colour=MUTED, italic=True)

    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    path = build()
    print(f"Wrote {path}")
