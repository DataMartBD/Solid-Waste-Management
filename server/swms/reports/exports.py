"""Server-side report export: CSV, Excel and PDF.

The browser can already build these from data it holds; doing it on the server
matters when a report is too large to ship to the client, when it has to be
identical for an audit, or when it is emailed by a scheduled job.

CSV values are neutralised against formula injection — a cell starting `=`, `+`,
`-`, `@`, tab or carriage return is prefixed with an apostrophe, so opening the
file in Excel cannot execute it. The frontend exporter does the same thing; both
paths need it because either can produce the file a user opens.
"""

from __future__ import annotations

import csv
import datetime as dt
import io
from typing import Any, Iterable, Sequence

from django.http import HttpResponse
from django.utils import timezone

#: Characters Excel and LibreOffice treat as the start of a formula.
_RISKY_PREFIXES = ("=", "+", "-", "@", "\t", "\r")

FORMATS = ("csv", "xlsx", "pdf")


class Column:
    """A report column: where the value comes from and what to call it."""

    __slots__ = ("key", "label", "getter", "align")

    def __init__(self, key: str, label: str, getter=None, align: str = "left"):
        self.key = key
        self.label = label
        self.getter = getter
        self.align = align

    def value(self, row: dict) -> Any:
        if self.getter is not None:
            return self.getter(row)
        return row.get(self.key)


def columns_from(spec: Sequence[tuple[str, str]]) -> list[Column]:
    return [Column(key, label) for key, label in spec]


def _cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, (dt.datetime,)):
        return timezone.localtime(value).strftime("%Y-%m-%d %H:%M")
    if isinstance(value, (dt.date,)):
        return value.isoformat()
    if isinstance(value, (list, tuple)):
        return ", ".join(_cell(item) for item in value)
    if isinstance(value, dict):
        return ", ".join(f"{k}: {_cell(v)}" for k, v in value.items())
    return str(value)


def _neutralise(text: str) -> str:
    return f"'{text}" if text[:1] in _RISKY_PREFIXES else text


def _filename(title: str, extension: str) -> str:
    stamp = timezone.localdate().isoformat()
    slug = "".join(ch if ch.isalnum() else "-" for ch in title.lower()).strip("-")
    while "--" in slug:
        slug = slug.replace("--", "-")
    return f"{slug or 'report'}-{stamp}.{extension}"


def to_csv(title: str, columns: Sequence[Column], rows: Iterable[dict]) -> HttpResponse:
    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator="\r\n")
    writer.writerow([column.label for column in columns])
    for row in rows:
        writer.writerow([_neutralise(_cell(column.value(row))) for column in columns])

    # UTF-8 BOM so Excel on Windows reads Bangla text correctly.
    payload = "﻿" + buffer.getvalue()
    response = HttpResponse(payload.encode("utf-8"), content_type="text/csv; charset=utf-8")
    response["Content-Disposition"] = f'attachment; filename="{_filename(title, "csv")}"'
    return response


def to_xlsx(
    title: str,
    columns: Sequence[Column],
    rows: Iterable[dict],
    *,
    subtitle: str = "",
) -> HttpResponse:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    sheet = workbook.active
    # Excel caps sheet names at 31 characters and forbids several punctuation marks.
    sheet.title = "".join(ch for ch in title if ch not in "[]:*?/\\")[:31] or "Report"

    heading = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="1F6F43")

    offset = 1
    if subtitle:
        sheet.cell(row=1, column=1, value=subtitle).font = Font(italic=True, size=9)
        offset = 2

    for index, column in enumerate(columns, start=1):
        cell = sheet.cell(row=offset, column=index, value=column.label)
        cell.font = heading
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center")

    widths = [len(column.label) for column in columns]
    for row_index, row in enumerate(rows, start=offset + 1):
        for index, column in enumerate(columns, start=1):
            raw = column.value(row)
            # Keep real numbers numeric so Excel can total them.
            value = raw if isinstance(raw, (int, float)) and not isinstance(raw, bool) else _cell(raw)
            sheet.cell(row=row_index, column=index, value=value)
            widths[index - 1] = max(widths[index - 1], len(str(value)))

    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = min(max(width + 2, 10), 48)
    sheet.freeze_panes = sheet.cell(row=offset + 1, column=1)

    stream = io.BytesIO()
    workbook.save(stream)
    response = HttpResponse(
        stream.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{_filename(title, "xlsx")}"'
    return response


def to_pdf(
    title: str,
    columns: Sequence[Column],
    rows: Iterable[dict],
    *,
    subtitle: str = "",
) -> HttpResponse:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    stream = io.BytesIO()
    document = SimpleDocTemplate(
        stream,
        pagesize=landscape(A4),
        leftMargin=12 * mm,
        rightMargin=12 * mm,
        topMargin=12 * mm,
        bottomMargin=12 * mm,
        title=title,
    )
    styles = getSampleStyleSheet()
    story = [Paragraph(title, styles["Title"])]
    if subtitle:
        story.append(Paragraph(subtitle, styles["Normal"]))
    story.append(Spacer(1, 6 * mm))

    # NOTE: the bundled Helvetica has no Bangla glyphs, so Bangla text renders as
    # boxes. Register a Unicode font (e.g. Noto Sans Bengali) here if PDFs need to
    # carry Bangla; CSV and Excel already handle it.
    data = [[column.label for column in columns]]
    body_style = styles["BodyText"]
    body_style.fontSize = 7.5
    body_style.leading = 9
    for row in rows:
        data.append([Paragraph(_cell(column.value(row)), body_style) for column in columns])

    table = Table(data, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F6F43")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, 0), 8),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#CBD5D0")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F3F7F4")]),
            ]
        )
    )
    story.append(table)
    document.build(story)

    response = HttpResponse(stream.getvalue(), content_type="application/pdf")
    response["Content-Disposition"] = f'attachment; filename="{_filename(title, "pdf")}"'
    return response


def export(
    fmt: str,
    title: str,
    columns: Sequence[Column],
    rows: Iterable[dict],
    *,
    subtitle: str = "",
) -> HttpResponse:
    """Dispatch to the requested format. `rows` is materialised once."""
    materialised = list(rows)
    if fmt == "xlsx":
        return to_xlsx(title, columns, materialised, subtitle=subtitle)
    if fmt == "pdf":
        return to_pdf(title, columns, materialised, subtitle=subtitle)
    return to_csv(title, columns, materialised)
