"""Bulk upload of holdings from Excel or CSV.

The client's minimum-effort setup: a city's existing register arrives as a
spreadsheet with one row per holding —

    Holding No · Name · Mobile · City · Ward · Area · Address · Route · GPS Location

— plus, optionally, Holding Type, Tier, Charge, Units and Agency. Every row is
validated first and the whole file is reported row by row, so an operator can
fix the sheet before anything is written. Importing then writes the valid rows
in one transaction; an invalid row is never half-imported.

How the columns land:

* **Name / Mobile** are the holding's owner — the holding is the unit billed
  and collected from.
* **Area** is the road (or mohalla) inside the ward. One that does not exist
  yet is created, and the preview says so, because a new city's roads are
  exactly what is being set up.
* **GPS Location** ("22.8456, 89.5403") pins and verifies the building — the
  register is the record of where it is — which also makes it routable.
* **Route** (id or name, in the same ward) puts the holding at the end of that
  route's walk. A route needs a located building, so it needs GPS.
* A row whose ward, area and holding number match an existing holding
  **updates** it rather than creating a duplicate.
"""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field

from django.db import transaction
from django.db.models import Max, Q
from django.utils import timezone

from swms.agencies.models import Agency
from swms.catalog.models import City, HoldingType, Road, Tier, Ward
from swms.fieldops.models import Route, RouteStop

from .models import Holding

MAX_ROWS = 5000
MAX_BYTES = 5 * 1024 * 1024

#: Column headings in the template, in order. Required ones first.
TEMPLATE_COLUMNS = [
    "Holding No", "Name", "Mobile", "City", "Ward", "Area", "Address", "Route",
    "GPS Location", "Holding Type", "Tier", "Charge", "Units", "Agency",
]
REQUIRED = {"holding_no", "name", "ward", "area"}

#: Normalised heading → field. Several spellings are accepted for each.
ALIASES = {
    "holdingno": "holding_no", "holding": "holding_no", "holdingnumber": "holding_no",
    "name": "name", "owner": "name", "ownername": "name",
    "mobile": "mobile", "phone": "mobile", "mobileno": "mobile", "phoneno": "mobile",
    "city": "city", "citycorporation": "city", "pourashova": "city",
    "ward": "ward", "wardno": "ward",
    "area": "area", "road": "area", "mohalla": "area", "block": "area",
    "address": "address",
    "route": "route",
    "gpslocation": "gps", "gps": "gps", "location": "gps", "latlng": "gps",
    "latitude": "lat", "lat": "lat", "longitude": "lng", "lng": "lng", "long": "lng",
    "holdingtype": "holding_type", "type": "holding_type",
    "tier": "tier", "charge": "charge", "monthlycharge": "charge",
    "units": "units", "unitstotal": "units", "flats": "units",
    "agency": "agency",
}

_BN_DIGITS = str.maketrans("০১২৩৪৫৬৭৮৯", "0123456789")
_PHONE = re.compile(r"^01[3-9]\d{8}$")


class ImportFileError(Exception):
    """The file itself cannot be read — wrong type, empty, too big, no headings."""


@dataclass
class Row:
    number: int  # the spreadsheet row, counting the heading as row 1
    data: dict
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    resolved: dict = field(default_factory=dict)
    action: str = "create"

    def as_dict(self) -> dict:
        return {
            "row": self.number,
            "holdingNo": self.data.get("holding_no", ""),
            "name": self.data.get("name", ""),
            "ward": self.resolved["ward"].id if "ward" in self.resolved else self.data.get("ward", ""),
            "area": self.data.get("area", ""),
            "action": self.action,
            "status": "error" if self.errors else "ok",
            "errors": self.errors,
            "warnings": self.warnings,
            "holding": self.resolved.get("holding_id"),
        }


# --------------------------------------------------------------------------- #
# Reading the file
# --------------------------------------------------------------------------- #


def _normalise_heading(value) -> str:
    return re.sub(r"[^a-z]", "", str(value or "").lower())


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip().translate(_BN_DIGITS)


def read_rows(upload) -> list[Row]:
    """Turn an uploaded .xlsx or .csv into `Row`s keyed by field name."""
    name = (getattr(upload, "name", "") or "").lower()
    if getattr(upload, "size", 0) > MAX_BYTES:
        raise ImportFileError(f"The file is larger than {MAX_BYTES // (1024 * 1024)} MB.")
    raw = upload.read()
    if not raw:
        raise ImportFileError("The file is empty.")

    if name.endswith(".csv"):
        text = raw.decode("utf-8-sig", errors="replace")
        table = list(csv.reader(io.StringIO(text)))
    elif name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        try:
            sheet = load_workbook(io.BytesIO(raw), read_only=True, data_only=True).worksheets[0]
        except Exception as exc:  # noqa: BLE001 — any unreadable workbook is the same answer
            raise ImportFileError("The Excel file could not be read.") from exc
        table = [list(row) for row in sheet.iter_rows(values_only=True)]
    else:
        raise ImportFileError("Upload an Excel (.xlsx) or CSV (.csv) file.")

    # The first non-empty row is the heading.
    while table and not any(_cell(v) for v in table[0]):
        table.pop(0)
    if not table:
        raise ImportFileError("The file has no rows.")
    fields = [ALIASES.get(_normalise_heading(h)) for h in table[0]]
    missing = REQUIRED - set(fields)
    if missing:
        wanted = {"holding_no": "Holding No", "name": "Name", "ward": "Ward", "area": "Area"}
        raise ImportFileError(
            "Missing column(s): " + ", ".join(wanted[m] for m in sorted(missing)) + "."
        )

    rows = []
    for offset, values in enumerate(table[1:], start=2):
        data = {}
        for key, value in zip(fields, values):
            if key and key not in data:
                data[key] = _cell(value)
        if not any(data.values()):
            continue  # a blank line in the sheet
        rows.append(Row(number=offset, data=data))
    if not rows:
        raise ImportFileError("The file has headings but no holdings.")
    if len(rows) > MAX_ROWS:
        raise ImportFileError(f"At most {MAX_ROWS} holdings can be uploaded at once.")
    return rows


# --------------------------------------------------------------------------- #
# Validating
# --------------------------------------------------------------------------- #


class _Lookups:
    """Everything a row is checked against, loaded once per file."""

    def __init__(self, user):
        self.user = user
        self.cities = list(City.objects.all())
        self.wards = list(Ward.objects.select_related("zone").all())
        self.holding_types = list(HoldingType.objects.filter(active=True))
        self.tiers = list(Tier.objects.all())
        self.agencies = list(Agency.objects.all())
        self.visible = user.visible_ward_ids() if user else None
        self.own_agency = user.visible_agency_id() if user else None

    @staticmethod
    def _match(rows, text, *attrs):
        needle = text.strip().lower()
        for row in rows:
            for attr in attrs:
                value = getattr(row, attr, None)
                if value and str(value).strip().lower() == needle:
                    return row
        return None

    def city(self, text):
        return self._match(self.cities, text, "id", "name", "name_bn")

    def ward(self, text, city):
        pool = [w for w in self.wards if city is None or w.zone.city_id == city.id]
        found = self._match(pool, text, "id", "name", "name_bn", "short_label")
        if found is None:
            # "14" or "Ward 14" — the number alone, within the city.
            digits = re.sub(r"\D", "", text)
            if digits:
                candidates = [
                    w for w in pool
                    if re.sub(r"\D", "", w.short_label).lstrip("0") == digits.lstrip("0")
                ]
                if len(candidates) == 1:
                    found = candidates[0]
        return found

    def holding_type(self, text):
        if not text:
            return self.holding_types[0] if self.holding_types else None
        return self._match(self.holding_types, text, "id", "label")

    def tier(self, text):
        return self._match(self.tiers, text, "id", "label")

    def agency(self, text):
        return self._match(self.agencies, text, "id", "short_code", "name")


def _gps(row: Row):
    lat_text, lng_text = row.data.get("lat", ""), row.data.get("lng", "")
    text = row.data.get("gps", "")
    if text and not (lat_text or lng_text):
        parts = [p for p in re.split(r"[,\s;]+", text) if p]
        if len(parts) != 2:
            row.errors.append("GPS Location must be 'latitude, longitude', e.g. 22.8456, 89.5403.")
            return None
        lat_text, lng_text = parts
    if not (lat_text or lng_text):
        return None
    try:
        lat, lng = float(lat_text), float(lng_text)
    except ValueError:
        row.errors.append("GPS Location must be two numbers.")
        return None
    if not (20.0 <= lat <= 27.0 and 87.0 <= lng <= 93.0):
        row.errors.append("GPS Location is outside Bangladesh.")
        return None
    return round(lat, 6), round(lng, 6)


def validate(rows: list[Row], user) -> list[Row]:
    look = _Lookups(user)
    seen: dict[tuple, int] = {}
    roads_to_create: set[tuple[str, str]] = set()

    for row in rows:
        d = row.data
        for key, label in (("holding_no", "Holding No"), ("name", "Name"),
                           ("ward", "Ward"), ("area", "Area")):
            if not d.get(key):
                row.errors.append(f"{label} is required.")

        city = None
        if d.get("city"):
            city = look.city(d["city"])
            if city is None:
                row.errors.append(f"Unknown city '{d['city']}'.")

        ward = None
        if d.get("ward") and not (d.get("city") and city is None):
            ward = look.ward(d["ward"], city)
            if ward is None:
                where = f" in {city.name}" if city else ""
                row.errors.append(f"Unknown ward '{d['ward']}'{where}.")
            elif look.visible is not None and ward.id not in look.visible:
                row.errors.append(f"Ward {ward.id} is outside your assigned area.")
                ward = None
        if ward is not None:
            row.resolved["ward"] = ward

        mobile = re.sub(r"[\s\-()]", "", d.get("mobile", ""))
        if mobile:
            if mobile.startswith("+880"):
                mobile = mobile[3:]
            elif mobile.startswith("880"):
                mobile = mobile[2:]
            if not _PHONE.match(mobile):
                row.errors.append(f"Mobile '{d['mobile']}' is not a valid number (01XXXXXXXXX).")
        row.resolved["mobile"] = mobile

        holding_type = look.holding_type(d.get("holding_type", ""))
        if holding_type is None:
            row.errors.append(f"Unknown holding type '{d.get('holding_type', '')}'.")
        row.resolved["holding_type"] = holding_type

        if d.get("tier"):
            tier = look.tier(d["tier"])
            if tier is None:
                row.errors.append(f"Unknown tier '{d['tier']}'.")
            row.resolved["tier"] = tier
        for key, label in (("charge", "Charge"), ("units", "Units")):
            if d.get(key):
                if not d[key].isdigit():
                    row.errors.append(f"{label} must be a whole number.")
                else:
                    row.resolved[key] = int(d[key])

        agency = None
        if d.get("agency"):
            agency = look.agency(d["agency"])
            if agency is None:
                row.errors.append(f"Unknown agency '{d['agency']}'.")
            elif look.own_agency is not None and agency.id != look.own_agency:
                row.errors.append("You can only upload holdings for your own agency.")
        row.resolved["agency_id"] = agency.id if agency else look.own_agency
        if agency and ward and agency.city_id and agency.city_id != ward.zone.city_id:
            row.errors.append(f"{agency.short_code} does not work in {ward.id}'s city.")

        gps = _gps(row)
        row.resolved["gps"] = gps

        if ward is None or not d.get("area") or not d.get("holding_no"):
            continue

        road = Road.objects.filter(ward=ward, name__iexact=d["area"]).first()
        row.resolved["road"] = road
        if road is None:
            if (ward.id, d["area"].lower()) not in roads_to_create:
                row.warnings.append(f"New area/road '{d['area']}' will be added to {ward.id}.")
                roads_to_create.add((ward.id, d["area"].lower()))

        key = (ward.id, d["area"].lower(), d["holding_no"].lower())
        if key in seen:
            row.errors.append(f"Same holding as row {seen[key]}.")
            continue
        seen[key] = row.number

        existing = (
            Holding.objects.filter(
                ward=ward, road=road, holding_no__iexact=d["holding_no"]
            ).first()
            if road is not None else None
        )
        if existing is not None:
            row.action = "update"
            row.resolved["holding_id"] = existing.id
            if look.own_agency is not None and existing.agency_id not in (None, look.own_agency):
                row.errors.append("That holding belongs to another agency.")

        if d.get("route"):
            route = Route.objects.filter(
                Q(pk__iexact=d["route"]) | Q(name__iexact=d["route"]), ward=ward, active=True
            ).first()
            if route is None:
                row.errors.append(f"Unknown route '{d['route']}' in {ward.id}.")
            elif gps is None and not (existing and existing.verified):
                row.errors.append("A route needs a GPS Location — only a located holding can be routed.")
            row.resolved["route"] = route

    return rows


# --------------------------------------------------------------------------- #
# Writing
# --------------------------------------------------------------------------- #


@transaction.atomic
def write(rows: list[Row], user) -> dict:
    """Create or update every valid row. Invalid rows are skipped, untouched."""
    created = updated = 0
    for row in rows:
        if row.errors:
            continue
        d, r = row.data, row.resolved
        ward = r["ward"]
        road = r.get("road") or Road.objects.filter(ward=ward, name__iexact=d["area"]).first()
        if road is None:
            road = Road.objects.create(ward=ward, name=d["area"].strip())

        holding = Holding.objects.filter(pk=r["holding_id"]).first() if r.get("holding_id") else None
        if holding is None:
            holding = Holding(ward=ward, road=road, holding_no=d["holding_no"].strip())
            created += 1
        else:
            updated += 1
        holding.owner_name = d["name"].strip()
        if r.get("mobile"):
            holding.owner_phone = r["mobile"]
        if d.get("address"):
            holding.address = d["address"].strip()
        holding.holding_type = r["holding_type"]
        if r.get("tier"):
            holding.tier = r["tier"]
        if "charge" in r:
            holding.charge = r["charge"]
        if "units" in r:
            holding.units_total = r["units"]
        if r.get("agency_id") and not holding.agency_id:
            holding.agency_id = r["agency_id"]
        if r.get("gps"):
            holding.lat, holding.lng = r["gps"]
            holding.placed_by_hand = True
            holding.verified = True
            holding.verified_at = timezone.now()
            holding.verified_by_user = user if getattr(user, "is_authenticated", False) else None
        holding.save()
        r["holding_id"] = holding.id

        route = r.get("route")
        if route is not None:
            stop = RouteStop.objects.filter(holding=holding).first()
            if stop is None or stop.route_id != route.id:
                last = RouteStop.objects.filter(route=route).aggregate(m=Max("seq"))["m"] or 0
                RouteStop.objects.update_or_create(
                    holding=holding, defaults={"route": route, "seq": last + 1}
                )
    return {"created": created, "updated": updated}


def run_import(upload, user, *, commit: bool) -> dict:
    rows = validate(read_rows(upload), user)
    valid = [row for row in rows if not row.errors]
    result = {"created": 0, "updated": 0}
    if commit and valid:
        result = write(rows, user)
    return {
        "dryRun": not commit,
        "rows": [row.as_dict() for row in rows],
        "summary": {
            "total": len(rows),
            "valid": len(valid),
            "invalid": len(rows) - len(valid),
            "toCreate": sum(1 for row in valid if row.action == "create"),
            "toUpdate": sum(1 for row in valid if row.action == "update"),
            **result,
        },
    }


def template_workbook() -> bytes:
    """An .xlsx with the headings, one example row and a notes sheet."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    book = Workbook()
    sheet = book.active
    sheet.title = "Holdings"
    sheet.append(TEMPLATE_COLUMNS)
    sheet.append([
        "142/B", "Abdul Karim", "01711000000", "KCC", "W-14", "KDA Avenue",
        "Holding 142/B, KDA Avenue", "", "22.8456, 89.5403", "", "", "", "4", "",
    ])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0F7B4E")
    for index, heading in enumerate(TEMPLATE_COLUMNS, start=1):
        sheet.column_dimensions[sheet.cell(row=1, column=index).column_letter].width = max(
            14, len(heading) + 4
        )

    notes = book.create_sheet("Notes")
    for line in [
        ["Column", "Required", "What to enter"],
        ["Holding No", "Yes", "The holding number, e.g. 142/B"],
        ["Name", "Yes", "Owner of the holding"],
        ["Mobile", "No", "Owner's mobile, 01XXXXXXXXX"],
        ["City", "No", "City code or name, e.g. KCC — helps find the ward"],
        ["Ward", "Yes", "Ward code (W-14), name, or number (14)"],
        ["Area", "Yes", "Road / mohalla in the ward; a new one is added"],
        ["Address", "No", "Full address; built from holding and road if blank"],
        ["Route", "No", "Route code or name in the same ward; needs GPS"],
        ["GPS Location", "No", "latitude, longitude — e.g. 22.8456, 89.5403"],
        ["Holding Type", "No", "From setup; the first type if blank"],
        ["Tier", "No", "Service tier code or name"],
        ["Charge", "No", "Monthly charge in BDT; blank = the tier's rate"],
        ["Units", "No", "Number of flats/units in the building"],
        ["Agency", "No", "Agency code (e.g. PCM); yours if you belong to one"],
    ]:
        notes.append(line)
    for cell in notes[1]:
        cell.font = Font(bold=True)
    notes.column_dimensions["A"].width = 16
    notes.column_dimensions["C"].width = 60

    out = io.BytesIO()
    book.save(out)
    return out.getvalue()
